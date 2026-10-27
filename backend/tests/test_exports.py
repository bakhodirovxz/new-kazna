import io
import time
import zipfile
from collections.abc import Callable, Iterator

import pytest
import redis
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import text

from app.core.config import get_settings
from app.core.db import get_engine
from app.main import create_app
from app.sections.deps import get_today
from app.sections.period import resolve_period
from app.sections.queries.cash import MonthPoint
from app.sections.queries.payments import PlanFactRow
from app.sections.schemas import BranchCashOut, CashKpisOut, DetailRowOut, ObzorExcelOut, period_out
from app.sections.xlsx import SHEETS, build_workbook
from tests.etl_helpers import TEST_AS_OF, new_source, purge_source
from tests.helpers import SeedUser, login
from tests.s1_helpers import Scenario, treasury_scenario


@pytest.fixture(autouse=True)
def _clean() -> None:
    r = redis.Redis.from_url(get_settings().redis_url)
    for key in [*r.scan_iter("ugk:*"), *r.scan_iter("rl:*")]:
        r.delete(key)


@pytest.fixture
def client_for() -> Iterator[Callable[[SeedUser], TestClient]]:
    opened: list[TestClient] = []

    def make(user: SeedUser) -> TestClient:
        app = create_app()
        app.dependency_overrides[get_today] = lambda: TEST_AS_OF
        client = TestClient(app, base_url="https://testserver")
        client.__enter__()
        opened.append(client)
        return login(client, user)

    yield make
    for client in opened:
        client.__exit__(None, None, None)


@pytest.fixture
def sc() -> Iterator[Scenario]:
    source = new_source()
    yield treasury_scenario(source, publish=True)
    purge_source(source)


def _last_export() -> dict[str, object]:
    with get_engine("migrator").connect() as conn:
        row = conn.execute(text("SELECT user_id, details FROM core.audit_log "
                                "WHERE action = 'data.exported' ORDER BY id DESC LIMIT 1")).one()
    return {"user_id": row.user_id, **row.details}


def test_xlsx_four_sheets_with_scenario_values(client_for: Callable[[SeedUser], TestClient],
                                               make_user: Callable[..., SeedUser],
                                               sc: Scenario) -> None:
    user = make_user("rukovodstvo")
    client = client_for(user)
    r = client.get("/api/v1/obzor/excel.xlsx", params={"per": "m08", "br": sc.branch_id})
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert r.headers["content-disposition"] == (
        f'attachment; filename="UGEKAZNA_cashflow_br{sc.branch_id}_m08_20260924.xlsx"')
    assert r.headers["cache-control"] == "no-store"
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == list(SHEETS)
    svod = [[c.value for c in row] for row in wb["Сводка"].iter_rows(min_row=3)]
    assert svod[0] == ["Филиал", "Начальный остаток", "Поступления", "Расходование",
                       "Чистый поток", "Конечный остаток", "Разница"]
    assert svod[1][1:] == [1000.0, 100.0, 40.0, 60.0, 1060.0, 0.0]
    assert svod[-1][0] == "Итого"
    plan = [[c.value for c in row] for row in wb["План-факт"].iter_rows(min_row=2)]
    assert plan[0][1:] == ["Операционная", 50.0, 40.0, 80.0, 10.0]
    months = [[c.value for c in row] for row in wb["По месяцам"].iter_rows(min_row=2)]
    assert months[7] == ["Август", 100.0, 40.0, 60.0] and len(months) == 9
    audit = _last_export()
    assert audit["user_id"] == user.id and audit["format"] == "xlsx" and audit["per"] == "m08"


def test_xlsx_within_ten_seconds_on_synthetic(client_for: Callable[[SeedUser], TestClient],
                                              make_user: Callable[..., SeedUser],
                                              published_synthetic: int) -> None:
    client = client_for(make_user("rukovodstvo"))
    started = time.perf_counter()
    r = client.get("/api/v1/obzor/excel.xlsx", params={"per": "ytd"})
    assert r.status_code == 200 and time.perf_counter() - started < 10


def test_xlsx_requires_export_permission(client_for: Callable[[SeedUser], TestClient],
                                         make_user: Callable[..., SeedUser]) -> None:
    auditor = client_for(make_user("vnutr_kontrol"))
    assert auditor.get("/api/v1/obzor/excel.xlsx").status_code == 403
    econ = client_for(make_user("economist_filiala", branch_ids=(3,)))
    assert econ.get("/api/v1/obzor/excel.xlsx", params={"br": 4}).status_code == 403


def test_formula_like_names_are_text() -> None:
    period = resolve_period("m08", TEST_AS_OF)
    kpis = CashKpisOut(opening=None, inflow=1, outflow=0, net=1, closing=None, diff=None,
                       opening_prev=None, inflow_prev=0, outflow_prev=0, net_prev=0,
                       closing_prev=None, diff_prev=None, has_data=True)
    branch = BranchCashOut(branch_id=9, name="+cmd|' /C calc'!A0", short_name="x", opening=None,
                           inflow=1, outflow=0, net=1, closing=None, diff=None)
    data = ObzorExcelOut(
        period=period_out(period), kpis=kpis, branches=[branch],
        total=branch.model_copy(update={"branch_id": 0, "name": "Итого"}),
        detail=[DetailRowOut(branch_id=9, branch_name="@SUM(A1)", direction="in", item_id=1,
                             item_name='=HYPERLINK("http://x","y")', activity="operating",
                             amount=1)])
    body = build_workbook(data, [PlanFactRow(1, "-2+3", "operating", 1, 1)],
                          [MonthPoint(202608, 1, 0, None, 0, 0)], "все филиалы")
    with zipfile.ZipFile(io.BytesIO(body)) as z:
        sheets = [n for n in z.namelist() if n.startswith("xl/worksheets/sheet")]
        assert sheets and all(b"<f>" not in z.read(n) for n in sheets)
    wb = load_workbook(io.BytesIO(body))
    assert wb["Детализация"]["C2"].value == '=HYPERLINK("http://x","y")'
    assert wb["Детализация"]["C2"].data_type == "s"
    assert wb["Сводка"]["A4"].value == "+cmd|' /C calc'!A0"
    assert wb["План-факт"]["A2"].value == "-2+3"


def test_export_log(client_for: Callable[[SeedUser], TestClient],
                    make_user: Callable[..., SeedUser]) -> None:
    user = make_user("rukovodstvo")
    client = client_for(user)
    body = {"kind": "csv", "section": "obzor", "page": "svod",
            "card": "Расходование по филиалам", "per": "m08", "rows": 5}
    assert client.post("/api/v1/exports/log", json=body).status_code == 204
    audit = _last_export()
    assert (audit["user_id"], audit["format"], audit["card"], audit["rows"]) == (
        user.id, "csv", "Расходование по филиалам", 5)
    assert client.post("/api/v1/exports/log", json={**body, "extra": 1}).status_code == 422
    assert client.post("/api/v1/exports/log", json={**body, "kind": "pdf"}).status_code == 422
    auditor = client_for(make_user("vnutr_kontrol"))
    assert auditor.post("/api/v1/exports/log", json=body).status_code == 403
    profile = client_for(make_user("profil_sluzhba"))
    assert profile.post("/api/v1/exports/log",
                        json={**body, "section": "kazna"}).status_code == 403
    econ = client_for(make_user("economist_filiala", branch_ids=(3,)))
    assert econ.post("/api/v1/exports/log", json={**body, "br": 4}).status_code == 403


def test_export_log_hardening(client_for: Callable[[SeedUser], TestClient],
                              make_user: Callable[..., SeedUser]) -> None:
    client = client_for(make_user("rukovodstvo"))
    body = {"kind": "csv", "section": "obzor", "page": "svod", "card": "ok", "per": "m08"}
    for card in ("a\nb", "x\x01y", "\x7f"):
        assert client.post("/api/v1/exports/log", json={**body, "card": card}).status_code == 422
    # page section'ga tegishli bo'lishi shart
    assert client.post("/api/v1/exports/log", json={**body, "page": "platezhi"}).status_code == 422
    assert client.post("/api/v1/exports/log", json=body).status_code == 204
    for _ in range(40):
        last = client.post("/api/v1/exports/log", json=body).status_code
    assert last == 429  # shortcut-limit 30/min


def test_xlsx_strips_illegal_control_chars() -> None:
    period = resolve_period("m08", TEST_AS_OF)
    kpis = CashKpisOut(opening=None, inflow=1, outflow=0, net=1, closing=None, diff=None,
                       opening_prev=None, inflow_prev=0, outflow_prev=0, net_prev=0,
                       closing_prev=None, diff_prev=None, has_data=True)
    branch = BranchCashOut(branch_id=9, name="Фил\x01иал\x0b", short_name="x", opening=None,
                           inflow=1, outflow=0, net=1, closing=None, diff=None)
    data = ObzorExcelOut(period=period_out(period), kpis=kpis, branches=[branch],
                         total=branch.model_copy(update={"branch_id": 0}), detail=[])
    body = build_workbook(data, [], [], "все филиалы")
    assert load_workbook(io.BytesIO(body))["Сводка"]["A4"].value == "Филиал"
