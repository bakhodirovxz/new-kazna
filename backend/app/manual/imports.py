"""«Ввод данных»: kanonik shablonlar va CSV/XLSX import. Fayl diskka yozilmaydi; hammasi yoki
hech narsa. Xato elementi: {"row": fayl qatori (1 = sarlavha), "field", "code"} -- QIYMATSIZ."""

from __future__ import annotations

import csv
import hashlib
import hmac
import io
import re
import threading
import unicodedata
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any
from urllib.parse import unquote

from sqlalchemy.orm import Session

from app.core import keystore
from app.core.errors import AppError, ValidationFailed
from app.etl.errors import FileFormatError
from app.etl.file_connector import _number
from app.manual.entities import FieldDef, ManualEntity
from app.manual.validate import check_items

CSV_TYPE = "text/csv"
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_ROWS = 5000
MAX_ERRORS = 200
# XLSX himoyasi (review C1): tana <= 1 MiB, lekin ochilgan hajm qismlar bo'yicha cheklanadi.
# 5 000 qator x 4-6 ustunli Excel faylida varaq ~1-1.5 MiB, sharedStrings < ~0.3 MiB,
# styles < ~60 KiB.
ZIP_MAX_MEMBERS = 100
ZIP_MAX_BYTES = 8 * 1024 * 1024  # barcha qismlar yig'indisi
CAP_SHEET = 4 * 1024 * 1024  # xl/worksheets/*.xml
CAP_SHARED = 2 * 1024 * 1024  # xl/sharedStrings.xml
CAP_STYLES = 512 * 1024  # xl/styles.xml
CAP_OTHER = 64 * 1024  # qolgan hamma qism (theme, workbook, docProps...)
MAX_RATIO = 100  # siqish nisbati (faqat 64 KiB dan katta qismlarda)
MAX_SHEET_ROWS = 20_000  # <row> elementlari soni
MAX_ROW_INDEX = 100_000  # qator raqami (bo'sh oraliqlarni to'ldirish chegarasi)
MAX_CELLS_PER_ROW = 256
_NUMERIC = frozenset({"money", "volume", "int", "decimal", "pct"})
_TEXT_KINDS = frozenset({"secret", "text"})
_BAD_PARTS = ("xl/externallinks/", "xl/activex/", "xl/embeddings/", "xl/oleobject")
_SHEET_PART = re.compile(r"xl/worksheets/[^/]+\.xml")
_CELL_REF = re.compile(r"[A-Za-z]+(\d+)$")
_xlsx_gate = threading.Semaphore(1)  # jarayonda bir vaqtda bitta XLSX tahlili
_KIND_RU = {"money": "сумма", "volume": "число", "int": "целое", "decimal": "число",
            "pct": "процент", "date": "дата ГГГГ-ММ-ДД", "month": "месяц ГГГГ-ММ",
            "text": "текст", "enum": "значение из списка", "ref": "код из справочника",
            "secret": "текст (скрыто)"}
Rows = list[tuple[int, list[Any]]]
Parsed = tuple[list[dict[str, str]], list[int]]  # qiymatlar, fayl qator raqamlari


def _unsupported() -> ValidationFailed:
    return ValidationFailed("unsupported_file")


def _invalid(rows: list[dict[str, Any]]) -> ValidationFailed:
    return ValidationFailed("import_invalid", rows=rows[:MAX_ERRORS],
                            fields=sorted({r["field"] for r in rows}))


class FormatErrors(ValidationFailed):
    """Katak formati xatolari + xatosiz qatorlar (spravochnik tekshiruvi ular ustida ham
    o'tkaziladi, shunda foydalanuvchi hamma xatoni bir yo'la ko'radi)."""

    def __init__(self, rows: list[dict[str, Any]], good: Parsed) -> None:
        super().__init__("import_invalid", rows=rows[:MAX_ERRORS],
                         fields=sorted({r["field"] for r in rows}))
        self.good = good


def with_reference_errors(db: Session, entity: ManualEntity, branch_id: int,
                          exc: FormatErrors) -> ValidationFailed:
    """Format xatolariga xatosiz qatorlarning mazmun/spravochnik xatolarini qo'shadi."""
    items, lines = exc.good
    rows = [*exc.params["rows"]]
    if items:
        _, found = check_items(db, entity, branch_id, items)
        rows += [{**e, "row": lines[e["row"] - 1]} for e in found]
    rows.sort(key=lambda e: (e["row"], e["field"]))
    return _invalid(rows)


def _safe(value: str) -> str:
    """Formula-injection: `=+-@` yoki TAB/CR bilan boshlangan matn `'` bilan."""
    return "'" + value if value[:1] in ("=", "+", "-", "@", "\t", "\r") else value


def header_label(f: FieldDef) -> str:
    """Shablon sarlavhasi (ruscha). Parser eski texnik kodni (`branch_code`) ham qabul qiladi."""
    if f.kind == "ref":
        return f"{f.title} (код)"
    return f"{f.title}, {f.unit}" if f.unit else f.title


def template_csv(entity: ManualEntity) -> bytes:
    out = io.StringIO()
    csv.writer(out, delimiter=";", lineterminator="\r\n").writerow(
        [_safe(header_label(f)) for f in entity.fields])
    return b"\xef\xbb\xbf" + out.getvalue().encode()


def template_xlsx(entity: ManualEntity) -> bytes:
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    data = wb.active
    assert data is not None
    data.title = "Данные"
    data.append([_safe(header_label(f)) for f in entity.fields])
    for i, f in enumerate(entity.fields, start=1):
        if f.kind in _TEXT_KINDS:  # Excel hisob raqamini songa aylantirmasin (review I2)
            data.column_dimensions[get_column_letter(i)].number_format = "@"
    info = wb.create_sheet("Описание")
    info.append(["Колонка", "Код поля", "Название", "Тип", "Обязательное", "Формат"])
    for f in entity.fields:
        fmt = "; ".join(c for c, _ in f.options) if f.options else f.unit
        info.append([_safe(header_label(f)), _safe(f.name), _safe(f.title), _KIND_RU[f.kind],
                     "да" if f.required else "нет", _safe(fmt)])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def safe_name(raw: str) -> str | None:
    """X-File-Name (faqat ko'rsatish uchun): NFC, Cc/Cf/Zl/Zp (bidi, nol kenglik, C1) olib
    tashlanadi, yo'ldan faqat nom, <= 200 belgi (review M3)."""
    text = unicodedata.normalize("NFC", unquote(raw))
    text = "".join(ch for ch in text if unicodedata.category(ch) not in ("Cc", "Cf", "Zl", "Zp"))
    return text.replace("\\", "/").rsplit("/", 1)[-1].strip()[:200] or None


def file_digest(body: bytes) -> bytes:
    """Faylning kalitli HMAC'i (audit/store): oddiy SHA-256 past entropiyali faylni (bitta hisob
    raqami) auditorga qo'pol kuch bilan ochtirardi (review I3). Kalit -- mavjud session_pepper."""
    return hmac.new(keystore.pepper(), b"import-file|" + body, hashlib.sha256).digest()


def _cell(f: FieldDef, value: Any, row: int, errors: list[dict[str, Any]]) -> str | None:
    """Katakni kanonik satrga o'giradi; bo'sh -> None."""
    bad = {"row": row, "field": f.name, "code": "invalid_value"}
    if value is None:
        return None
    if isinstance(value, bool):
        errors.append(bad)
        return None
    if f.kind in _TEXT_KINDS and isinstance(value, float | int | Decimal):
        # 20 xonali hisob raqami float orqali jimgina buziladi (review I2)
        errors.append({**bad, "code": "text_required"})
        return None
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float | int | Decimal):
        try:  # float -> kanonik o'nlik satr (1e-05 ham)
            value = format(Decimal(repr(value)), "f")
        except InvalidOperation:
            errors.append(bad)
            return None
    text = str(value).strip()
    if text == "":
        return None
    if f.kind in _NUMERIC:
        try:
            return _number(text, "")
        except FileFormatError:
            errors.append(bad)
            return None
    return text


def _build(entity: ManualEntity, header: list[Any], rows: Rows) -> Parsed:
    known = {k.casefold(): f.name for f in entity.fields for k in (f.name, header_label(f))}
    order = [known.get(str(h or "").strip().casefold(), "") for h in header]
    if sorted(order) != sorted(f.name for f in entity.fields):
        raise _invalid([{"row": 1, "field": "", "code": "bad_header"}])
    by_name = {f.name: f for f in entity.fields}
    if len(rows) > MAX_ROWS:
        raise _invalid([{"row": rows[MAX_ROWS][0], "field": "", "code": "too_many_rows"}])
    if not rows:
        raise _invalid([{"row": 2, "field": "", "code": "required"}])
    items: list[dict[str, str]] = []
    lines: list[int] = []
    errors: list[dict[str, Any]] = []
    items_ok: list[dict[str, str]] = []
    lines_ok: list[int] = []
    for line, raw in rows:
        if len(raw) != len(order):
            errors.append({"row": line, "field": "", "code": "invalid_value"})
            continue
        item: dict[str, str] = {}
        before = len(errors)
        for name, value in zip(order, raw, strict=True):
            v = _cell(by_name[name], value, line, errors)
            if v is not None:
                item[name] = v
        items.append(item)
        lines.append(line)
        if len(errors) == before:
            items_ok.append(item)
            lines_ok.append(line)
    if errors:
        raise FormatErrors(errors, (items_ok, lines_ok))
    return items, lines


def _parse_csv(entity: ManualEntity, body: bytes) -> Parsed:
    try:
        reader = csv.reader(io.StringIO(body.decode("utf-8-sig")), delimiter=";")
        lines: Rows = []
        prev = 0
        for raw in reader:  # yozuv boshlanish qatori (bo'sh va ko'p qatorli kataklar bilan)
            start, prev = prev + 1, reader.line_num
            if any(c.strip() for c in raw):
                lines.append((start, raw))
    except (UnicodeDecodeError, csv.Error):
        raise _unsupported() from None
    if not lines:
        raise _invalid([{"row": 1, "field": "", "code": "bad_header"}])
    return _build(entity, lines[0][1], lines[1:])


def _part_name(name: str) -> str:
    return name.replace("\\", "/").lstrip("/").lower()


def _check_zip(body: bytes) -> list[zipfile.ZipInfo]:
    try:
        with zipfile.ZipFile(io.BytesIO(body)) as zf:
            infos = zf.infolist()
    except zipfile.BadZipFile:
        raise _unsupported() from None
    if len(infos) > ZIP_MAX_MEMBERS or sum(i.file_size for i in infos) > ZIP_MAX_BYTES:
        raise _unsupported()
    for i in infos:
        name = _part_name(i.filename)
        if name.endswith("vbaproject.bin") or name.startswith(_BAD_PARTS):
            raise _unsupported()
        if _SHEET_PART.fullmatch(name):
            cap = CAP_SHEET
        elif name == "xl/sharedstrings.xml":
            cap = CAP_SHARED
        elif name == "xl/styles.xml":
            cap = CAP_STYLES
        else:
            cap = CAP_OTHER
        ratio_bad = i.file_size > CAP_OTHER and i.compress_size * MAX_RATIO < i.file_size
        if i.file_size > cap or ratio_bad:
            raise _unsupported()
    return infos


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _scan_sheet(stream: Any) -> None:
    """Varaq XML'ini oqim bilan tekshiradi (xotira o'smaydi): keshsiz formula, juda keng/uzun
    qatorlar, haddan ko'p qator. openpyxl'dan OLDIN -- uning xarajati shu chegaralar bilan
    cheklanadi."""
    from defusedxml.ElementTree import iterparse  # type: ignore[import-untyped]

    rows = filled = cells = 0
    row_has_value = False
    last_row = 1
    for _, el in iterparse(stream, events=("end",)):
        tag = _local(el.tag)
        if tag == "c":
            ref = _CELL_REF.match(el.get("r", ""))
            last_row = int(ref.group(1)) if ref else last_row
            kids = {_local(k.tag): k for k in el}
            has_value = (kids["v"].text or "") != "" if "v" in kids else "is" in kids
            if "f" in kids and not has_value:
                raise _invalid([{"row": last_row, "field": "", "code": "formula_no_value"}])
            cells += 1
            row_has_value = row_has_value or has_value
            if cells > MAX_CELLS_PER_ROW:
                raise _invalid([{"row": last_row, "field": "", "code": "bad_header"}])
            el.clear()
        elif tag == "row":
            rows += 1
            idx = int(el.get("r", "0") or 0)
            filled += row_has_value
            if rows > MAX_SHEET_ROWS or idx > MAX_ROW_INDEX or filled > MAX_ROWS + 1:
                raise _invalid([{"row": idx or last_row, "field": "", "code": "too_many_rows"}])
            cells, row_has_value = 0, False
            el.clear()


def _xlsx_rows(body: bytes, width: int) -> list[list[Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        ws.reset_dimensions()  # <dimension> ga ishonilmaydi (review I1)
        return [[*r] for r in ws.iter_rows(max_col=width + 1, values_only=True)]
    finally:
        wb.close()


def _parse_xlsx(entity: ManualEntity, body: bytes) -> Parsed:
    import openpyxl

    if not getattr(openpyxl, "DEFUSEDXML", False):  # defusedxml yo'q bo'lsa fail-closed (M1)
        raise _unsupported()
    infos = _check_zip(body)
    width = len(entity.fields)
    if not _xlsx_gate.acquire(blocking=False):
        raise AppError("import_busy", status_code=429, headers={"Retry-After": "5"})
    try:
        with zipfile.ZipFile(io.BytesIO(body)) as zf:
            for i in infos:
                if _SHEET_PART.fullmatch(_part_name(i.filename)):
                    with zf.open(i) as stream:
                        _scan_sheet(stream)
        grid = _xlsx_rows(body, width)
    except AppError:
        raise
    except Exception:  # buzuq/zararli XML (entity kengayishi ham); sabab oshkor emas
        raise _unsupported() from None
    finally:
        _xlsx_gate.release()
    if not grid:
        raise _invalid([{"row": 1, "field": "", "code": "bad_header"}])
    header = [*grid[0]]
    while header and header[-1] is None:
        header.pop()
    rows: Rows = []
    for line, r in enumerate(grid[1:], start=2):
        padded = [*r, *([None] * (width + 1 - len(r)))]
        if any(c is not None and str(c).strip() for c in padded):
            rows.append((line, padded))
    if len(header) == width and any(r[width] is not None for _, r in rows):
        raise _invalid([{"row": line, "field": "", "code": "invalid_value"}
                        for line, r in rows if r[width] is not None])
    return _build(entity, header, [(line, r[:width]) for line, r in rows])


def parse_upload(entity: ManualEntity, body: bytes, content_type: str) -> Parsed:
    if content_type == CSV_TYPE:
        return _parse_csv(entity, body)
    if content_type == XLSX_TYPE:
        return _parse_xlsx(entity, body)
    raise _unsupported()
