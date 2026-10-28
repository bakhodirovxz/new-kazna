"""«Ввод данных»: shablonlar va CSV/XLSX import (hammasi yoki hech narsa, ZIP/XML bomb, makros)."""

from __future__ import annotations

import hashlib
import io
import json
import resource
import zipfile
from collections.abc import Callable, Iterator
from datetime import datetime
from urllib.parse import quote

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.core.errors import ValidationFailed
from app.main import create_app
from app.manual import imports
from app.manual.entities import ENTITIES
from tests.helpers import SeedUser, login
from tests.manual_helpers import audit_rows, branch_code, migrator_sql, seed_dwh_bank

pytestmark = pytest.mark.usefixtures("clean_manual")
MakeUser = Callable[..., SeedUser]
BANK = "BA-t7bank0000000001"
BALANCE = ENTITIES["kazna.balance"]
CSV = "text/csv"
XLSX = imports.XLSX_TYPE
URL = "/api/v1/data/entities/kazna.balance/import?branch_id=1"
HEAD = "branch_code;bank_code;date;balance"  # eski texnik kodlar ham qabul qilinadi
RU_HEAD = ["Филиал (код)", "Счёт (код)", "Дата", "Остаток, млн сум"]


@pytest.fixture
def client_for() -> Iterator[Callable[[SeedUser], TestClient]]:
    opened: list[TestClient] = []

    def make(user: SeedUser) -> TestClient:
        client = TestClient(create_app(), base_url="https://testserver")
        client.__enter__()
        opened.append(client)
        return login(client, user)

    seed_dwh_bank(BANK)
    yield make
    for client in opened:
        client.__exit__(None, None, None)


def _row(day: str = "2026-09-30", value: str = "1500.00") -> str:
    return f"{branch_code(1)};{BANK};{day};{value}"


def _csv(*rows: str, head: str = HEAD, bom: bool = True) -> bytes:
    return (("﻿" if bom else "") + "\r\n".join([head, *rows]) + "\r\n").encode()


def _post(client: TestClient, body: bytes, ctype: str = CSV, name: str = "ostatki.csv",
          url: str = URL):  # type: ignore[no-untyped-def]
    return client.post(url, content=body, headers={
        "Content-Type": ctype, "X-File-Name": quote(name)})


def _xlsx(rows: list[list[object]], header: list[str] | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(header or [f.name for f in BALANCE.fields])
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _batches() -> int:
    return int(migrator_sql("SELECT count(*) FROM core.manual_batch")[0][0])


def _entries() -> int:
    return int(migrator_sql("SELECT count(*) FROM core.manual_entry")[0][0])


def _codes(r) -> list[tuple[int, str, str]]:  # type: ignore[no-untyped-def]
    return [(x["row"], x["field"], x["code"]) for x in r.json()["error"]["params"]["rows"]]


def test_csv_import_creates_draft(client_for, make_user: MakeUser) -> None:  # type: ignore[no-untyped-def]
    maker = client_for(make_user("spec_kazna"))
    body = _csv(_row(), _row("2026-08-29", "20.50"))
    r = _post(maker, body)
    assert r.status_code == 201, r.text
    out = r.json()
    assert out["status"] == "draft" and out["rows"] == 2 and out["origin"] == "import"
    name, sha = migrator_sql("SELECT file_name, file_sha256 FROM core.manual_batch")[0]
    assert (name, bytes(sha)) == ("ostatki.csv", imports.file_digest(body))
    assert bytes(sha) != hashlib.sha256(body).digest()  # I3: kalitli HMAC, oddiy SHA-256 emas
    ev = [a for a in audit_rows("data.imported") if a["result"] == "ok"]
    assert len(ev) == 1
    assert ev[0]["details"]["hmac"] == imports.file_digest(body).hex()
    assert hashlib.sha256(body).hexdigest() not in json.dumps(ev[0]["details"])
    assert ev[0]["details"]["rows"] == 2 and ev[0]["details"]["file_name"] == "ostatki.csv"
    assert "1500" not in json.dumps(ev[0]["details"])  # mazmun auditda yo'q


def test_bom_and_comma_decimal(client_for, make_user: MakeUser) -> None:  # type: ignore[no-untyped-def]
    maker = client_for(make_user("spec_kazna"))
    r = _post(maker, _csv(_row(value="1234,50"), bom=False))
    assert r.status_code == 201, r.text
    assert _post(maker, _csv(_row("2026-08-29", "7,25"), bom=True)).status_code == 201
    vals = migrator_sql("SELECT payload ->> 'balance' FROM core.manual_entry ORDER BY id")
    assert [v[0] for v in vals] == ["1234.50", "7.25"]


def test_thousands_separator_rejected(client_for, make_user: MakeUser) -> None:  # type: ignore[no-untyped-def]
    maker = client_for(make_user("spec_kazna"))
    r = _post(maker, _csv(_row(value="1 234,50")))
    assert r.status_code == 422 and r.json()["error"]["code"] == "import_invalid"
    assert _codes(r) == [(2, "balance", "invalid_value")]
    assert _batches() == 0


def test_bad_header(client_for, make_user: MakeUser) -> None:  # type: ignore[no-untyped-def]
    maker = client_for(make_user("spec_kazna"))
    for head in (HEAD + ";extra", "branch_code;bank_code;date"):
        r = _post(maker, _csv(_row(), head=head))
        assert r.status_code == 422 and _codes(r) == [(1, "", "bad_header")]
    assert _batches() == 0


def test_all_or_nothing(client_for, make_user: MakeUser) -> None:  # type: ignore[no-untyped-def]
    maker = client_for(make_user("spec_kazna"))
    r = _post(maker, _csv(_row(), _row("2026-08-29", "1.005"), _row("2026-07-28")))
    assert r.status_code == 422 and r.json()["error"]["code"] == "import_invalid"
    assert _codes(r) == [(3, "balance", "invalid_value")]  # fayl qatori (1 = sarlavha)
    assert _batches() == 0 and _entries() == 0


def test_5001_rows_rejected(client_for, make_user: MakeUser) -> None:  # type: ignore[no-untyped-def]
    maker = client_for(make_user("spec_kazna"))
    r = _post(maker, _csv(*["a;b;c;d"] * 5001))
    assert r.status_code == 422 and _codes(r) == [(5002, "", "too_many_rows")]
    assert _batches() == 0


def test_xlsx_roundtrip_template(client_for, make_user: MakeUser) -> None:  # type: ignore[no-untyped-def]
    maker = client_for(make_user("spec_kazna"))
    tpl = maker.get("/api/v1/data/entities/kazna.balance/template.xlsx")
    assert tpl.status_code == 200 and "attachment" in tpl.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(tpl.content))
    assert wb.sheetnames == ["Данные", "Описание"]
    assert [c.value for c in wb["Данные"][1]] == RU_HEAD
    ws = wb["Данные"]
    ws.append([branch_code(1), BANK, "2026-09-30", 1500.5])
    ws.append([branch_code(1), BANK, datetime(2026, 8, 29), 3])  # noqa: DTZ001 -- openpyxl tz'siz sana yozadi
    buf = io.BytesIO()
    wb.save(buf)
    r = _post(maker, buf.getvalue(), XLSX, "t.xlsx")
    assert r.status_code == 201, r.text
    vals = migrator_sql("SELECT payload ->> 'date', payload ->> 'balance' "
                        "FROM core.manual_entry ORDER BY id")
    assert vals == [("2026-09-30", "1500.50"), ("2026-08-29", "3.00")]


def test_xlsx_zip_bomb_rejected(client_for, make_user: MakeUser) -> None:  # type: ignore[no-untyped-def]
    maker = client_for(make_user("spec_kazna"))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("xl/worksheets/sheet1.xml", b"\0" * (21 * 1024 * 1024))
    assert len(buf.getvalue()) < 1_048_576
    r = _post(maker, buf.getvalue(), XLSX)
    assert r.status_code == 422 and r.json()["error"]["code"] == "unsupported_file"
    many = io.BytesIO()
    with zipfile.ZipFile(many, "w") as zf:
        for i in range(101):
            zf.writestr(f"f{i}", b"x")
    assert _post(maker, many.getvalue(), XLSX).json()["error"]["code"] == "unsupported_file"
    assert _batches() == 0


@pytest.mark.parametrize("member", ["xl/vbaProject.bin", "xl/externalLinks/externalLink1.xml"])
def test_xlsx_macro_rejected(client_for, make_user: MakeUser, member: str) -> None:  # type: ignore[no-untyped-def]
    maker = client_for(make_user("spec_kazna"))
    src = _xlsx([[branch_code(1), BANK, "2026-09-30", 1]])
    buf = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(src)) as zin, zipfile.ZipFile(buf, "w") as zout:
        for n in zin.namelist():
            zout.writestr(n, zin.read(n))
        zout.writestr(member, b"x")
    r = _post(maker, buf.getvalue(), XLSX)
    assert r.status_code == 422 and r.json()["error"]["code"] == "unsupported_file"
    assert _batches() == 0


def test_billion_laughs_rejected(client_for, make_user: MakeUser) -> None:  # type: ignore[no-untyped-def]
    maker = client_for(make_user("spec_kazna"))
    src = _xlsx([[branch_code(1), BANK, "2026-09-30", 1]])
    lols = "".join(
        f'<!ENTITY lol{i} "{f"&lol{i - 1};" * 10}">' for i in range(1, 10))
    bomb = ('<?xml version="1.0"?><!DOCTYPE sst [<!ENTITY lol0 "lollollollol">' + lols + "]>"
            '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="1" '
            'uniqueCount="1"><si><t>&lol9;</t></si></sst>').encode()
    buf = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(src)) as zin, zipfile.ZipFile(buf, "w") as zout:
        for n in zin.namelist():
            data = zin.read(n)
            if n == "[Content_Types].xml":  # openpyxl qismni faqat manifest orqali topadi
                data = data.replace(
                    b"</Types>",
                    b'<Override PartName="/xl/sharedStrings.xml" ContentType="application/'
                    b'vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>')
            zout.writestr(n, data)
        zout.writestr("xl/sharedStrings.xml", bomb)
    before = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss  # KiB
    r = _post(maker, buf.getvalue(), XLSX)
    grown = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss - before
    assert r.status_code == 422 and r.json()["error"]["code"] == "unsupported_file"
    assert grown < 200 * 1024
    assert _batches() == 0


def test_formula_without_cache_rejected(client_for, make_user: MakeUser) -> None:  # type: ignore[no-untyped-def]
    maker = client_for(make_user("spec_kazna"))
    body = _xlsx([[branch_code(1), BANK, "2026-09-30", "=1+1"]])  # openpyxl keshni yozmaydi
    r = _post(maker, body, XLSX)
    assert r.status_code == 422 and r.json()["error"]["code"] == "import_invalid"
    assert _codes(r)[0][2] == "formula_no_value"
    assert _batches() == 0


def test_wrong_content_type_422(client_for, make_user: MakeUser) -> None:  # type: ignore[no-untyped-def]
    maker = client_for(make_user("spec_kazna"))
    for ctype in ("text/plain", "application/octet-stream"):
        r = _post(maker, _csv(_row()), ctype)
        assert r.status_code == 422 and r.json()["error"]["code"] == "unsupported_file"
    assert _post(maker, _csv(_row()), "application/json").status_code == 422
    assert _batches() == 0


def test_body_over_1mib_413(client_for, make_user: MakeUser) -> None:  # type: ignore[no-untyped-def]
    maker = client_for(make_user("spec_kazna"))
    assert _post(maker, b"x" * (1_048_576 + 1)).status_code == 413


def test_import_rate_limited(client_for, make_user: MakeUser) -> None:  # type: ignore[no-untyped-def]
    maker = client_for(make_user("spec_kazna"))
    codes = [_post(maker, _csv(_row(), head="x")).status_code for _ in range(11)]
    assert codes[:10] == [422] * 10 and codes[10] == 429


def test_branch_and_permission_scope(client_for, make_user: MakeUser) -> None:  # type: ignore[no-untyped-def]
    econ = client_for(make_user("economist_filiala", branch_ids=(2,)))
    r = _post(econ, _csv(_row()), url=URL.replace("branch_id=1", "branch_id=3"))
    assert r.status_code == 403 and r.json()["error"]["code"] == "scope_violation"
    boss = client_for(make_user("nach_otdela"))  # faqat approve -- write yo'q
    assert _post(boss, _csv(_row())).status_code == 403
    assert boss.get("/api/v1/data/entities/kazna.balance/template.csv").status_code == 403
    assert _batches() == 0


def test_template_formula_injection_safe() -> None:
    assert imports._safe("=cmd|x") == "'=cmd|x"
    for ch in "=+-@":
        assert imports._safe(ch + "a") == "'" + ch + "a"
    assert imports._safe("balance") == "balance"
    raw = imports.template_csv(BALANCE)
    assert raw.startswith(b"\xef\xbb\xbf")
    assert raw.decode("utf-8-sig").splitlines() == [";".join(RU_HEAD)]
    ws = load_workbook(io.BytesIO(imports.template_xlsx(BALANCE)))["Описание"]
    assert [c.value for c in ws[1]][:3] == ["Колонка", "Код поля", "Название"]
    assert [row[1].value for row in ws.iter_rows(min_row=2)] == [f.name for f in BALANCE.fields]
    assert not any(str(c.value).startswith(("=", "+", "@")) for row in ws.iter_rows() for c in row)


def test_template_csv_endpoint(client_for, make_user: MakeUser) -> None:  # type: ignore[no-untyped-def]
    maker = client_for(make_user("spec_kazna"))
    r = maker.get("/api/v1/data/entities/kazna.balance/template.csv")
    assert r.status_code == 200 and r.content.startswith(b"\xef\xbb\xbf")
    assert "attachment" in r.headers["content-disposition"]


# --- review round 1 ---------------------------------------------------------------------------

ACCOUNT = ENTITIES["kazna.bank_account"]
SHEET = "xl/worksheets/sheet1.xml"
PARSE_SCRIPT = """
import json, sys, time


def hwm():  # VmHWM exec'da nolga tushadi (ru_maxrss ota-ona cho'qqisini meros qiladi)
    for line in open('/proc/self/status'):
        if line.startswith('VmHWM'):
            return int(line.split()[1]) / 1024


from app.core.errors import AppError
from app.manual import imports
from app.manual.entities import ENTITIES
body = open(sys.argv[1], 'rb').read()
base = hwm()
t = time.time()
try:
    imports.parse_upload(ENTITIES['kazna.balance'], body, imports.XLSX_TYPE)
    out = 'ok'
except AppError as exc:
    out = exc.code
print(json.dumps({'out': out, 'secs': time.time() - t, 'peak': hwm(),
                  'mib': hwm() - base}))
"""
MIB = 1024 * 1024


def _patched(src: bytes, parts: dict[str, bytes | None]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(src)) as zin, zipfile.ZipFile(
            buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in zin.namelist():
            if n not in parts:
                zout.writestr(n, zin.read(n))
        for n, data in parts.items():
            if data is not None:
                zout.writestr(n, data)
    return buf.getvalue()


def _valid() -> bytes:
    return _xlsx([[branch_code(1), BANK, "2026-09-30", 1]])


def _sheet(rows_xml: str) -> bytes:
    return ('<?xml version="1.0"?><worksheet xmlns="http://schemas.openxmlformats.org/'
            f'spreadsheetml/2006/main"><sheetData>{rows_xml}</sheetData></worksheet>').encode()


NS = b'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
STYLES_HEAD = b"<styleSheet " + NS + b"><cellXfs>"
STYLES_TAIL = b"</cellXfs></styleSheet>"
SST_HEAD = b"<sst " + NS + b">"

CRAFTED = {
    "styles_19mib": lambda: {
        "xl/styles.xml": STYLES_HEAD + b"<xf/>" * (19 * MIB // 4) + STYLES_TAIL},
    "styles_400k": lambda: {"xl/styles.xml": STYLES_HEAD + "".join(
        f'<xf numFmtId="{i}" fontId="{i}"/>' for i in range(12_000)).encode() + STYLES_TAIL},
    "sst_19mib": lambda: {
        "xl/sharedStrings.xml": SST_HEAD + b"<si/>" * (19 * MIB // 4) + b"</sst>"},
    "sst_1_9mib": lambda: {"xl/sharedStrings.xml": SST_HEAD + "".join(
        f"<si><t>s{i}</t></si>" for i in range(80_000)).encode() + b"</sst>"},
    "wide_row_19mib": lambda: {SHEET: _sheet('<row r="1">' + "<c/>" * (19 * MIB // 4) + "</row>")},
    "wide_row_3mib": lambda: {SHEET: _sheet('<row r="1">' + "".join(
        f'<c r="A1"><v>{i}</v></c>' for i in range(100_000)) + "</row>")},
    "many_rows_3mib": lambda: {SHEET: _sheet(
        "".join(f'<row r="{i}"/>' for i in range(1, 170_000)))},
}


@pytest.mark.parametrize("name", sorted(CRAFTED))
def test_crafted_xlsx_is_cheap(name: str, tmp_path) -> None:  # type: ignore[no-untyped-def]
    import os
    import subprocess
    import sys

    path = tmp_path / "f.xlsx"
    path.write_bytes(_patched(_valid(), CRAFTED[name]()))
    assert path.stat().st_size < 1_048_576
    res = subprocess.run([sys.executable, "-c", PARSE_SCRIPT, str(path)], capture_output=True,
                         text=True, check=True, cwd="/app",
                         env={**os.environ, "PYTHONPATH": "/app"})
    out = json.loads(res.stdout.strip().splitlines()[-1])
    print(name, out)  # noqa: T201 -- o'lchangan son hisobot uchun
    assert out["secs"] < 5 and out["mib"] < 150
    assert out["out"] in {"ok", "unsupported_file", "import_invalid"}


def test_xlsx_dimension_is_not_trusted() -> None:
    rows = [[branch_code(1), BANK, f"2026-{m:02d}-28", m] for m in range(1, 10)]
    src = _xlsx(rows + [[branch_code(1), BANK, f"2025-{m:02d}-28", m] for m in range(1, 10)])
    with zipfile.ZipFile(io.BytesIO(src)) as zf:
        sheet = zf.read(SHEET)
    assert b'ref="A1:D19"' in sheet
    body = _patched(src, {SHEET: sheet.replace(b'ref="A1:D19"', b'ref="A1:D2"')})
    items, lines = imports.parse_upload(BALANCE, body, XLSX)
    assert len(items) == 18 and lines[-1] == 19


def test_xlsx_rows_after_a_gap_are_not_dropped() -> None:
    wb = Workbook()
    ws = wb.active
    ws.append([f.name for f in BALANCE.fields])
    ws.append([branch_code(1), BANK, "2026-01-28", 1])
    for col, value in enumerate([branch_code(1), BANK, "2026-02-28", 2], start=1):
        ws.cell(row=7012, column=col, value=value)
    buf = io.BytesIO()
    wb.save(buf)
    items, lines = imports.parse_upload(BALANCE, buf.getvalue(), XLSX)
    assert len(items) == 2 and lines == [2, 7012]


def test_xlsx_over_5000_rows_is_an_error_not_a_truncation(client_for, make_user: MakeUser) -> None:  # type: ignore[no-untyped-def]
    maker = client_for(make_user("spec_kazna"))
    body = _xlsx([[branch_code(1), BANK, "2026-01-28", 1]] * 5001)
    r = _post(maker, body, XLSX)
    assert r.status_code == 422 and _codes(r)[0][2] == "too_many_rows"
    assert _batches() == 0


def test_numeric_cell_for_text_or_secret_is_rejected() -> None:
    num = 20208000900123456789
    header = [f.name for f in ACCOUNT.fields]
    body = _xlsx([[branch_code(1), "Асакабанк", num, "UZS"]], header=header)
    with pytest.raises(ValidationFailed) as exc:
        imports.parse_upload(ACCOUNT, body, XLSX)
    assert exc.value.params["rows"] == [
        {"row": 2, "field": "account_number", "code": "text_required"}]
    ok = _xlsx([[branch_code(1), "Асакабанк", str(num), "UZS"]], header=header)
    assert imports.parse_upload(ACCOUNT, ok, XLSX)[0][0]["account_number"] == str(num)
    ws = load_workbook(io.BytesIO(imports.template_xlsx(ACCOUNT)))["Данные"]
    letter = "ABCDEFG"[header.index("account_number")]
    assert ws.column_dimensions[letter].number_format == "@"


def test_rejected_uploads_are_audited(client_for, make_user: MakeUser) -> None:  # type: ignore[no-untyped-def]
    maker = client_for(make_user("spec_kazna"))
    seen = len(audit_rows("data.import.rejected"))
    macro = _patched(_valid(), {"xl/vbaProject.bin": b"x"})
    assert _post(maker, macro, XLSX, "evil.xlsx").status_code == 422
    assert _post(maker, _csv(_row(value="1.005")), CSV, "bad.csv").status_code == 422
    assert _post(maker, b"x", "text/plain", "a.txt").status_code == 422
    rows = audit_rows("data.import.rejected")[seen:]
    assert [r["details"]["error"] for r in rows] == [
        "unsupported_file", "import_invalid", "unsupported_file"]
    d = rows[0]["details"]
    assert d["entity"] == "kazna.balance" and d["file_name"] == "evil.xlsx"
    assert d["size"] == len(macro) and d["hmac"] == imports.file_digest(macro).hex()
    assert "1.005" not in json.dumps([r["details"] for r in rows])
    assert all(r["result"] == "fail" for r in rows)


def test_error_rows_use_physical_file_lines(client_for, make_user: MakeUser) -> None:  # type: ignore[no-untyped-def]
    maker = client_for(make_user("spec_kazna"))
    bad_ref = f"{branch_code(1)};BA-nonexistent;2026-08-29;1"
    r = _post(maker, _csv(_row(), "", bad_ref))  # 1 sarlavha, 2 qator, 3 bo'sh, 4 xato
    assert r.status_code == 422 and _codes(r) == [(4, "bank_code", "reference_not_found")]


def test_collision_reports_the_row(client_for, make_user: MakeUser) -> None:  # type: ignore[no-untyped-def]
    maker = client_for(make_user("spec_kazna"))
    item = {"values": {"branch_code": branch_code(1), "bank_code": BANK, "date": "2026-09-30",
                       "balance": "5.00"}}
    r = maker.post("/api/v1/data/batches", json={
        "entity": "kazna.balance", "branch_id": 1, "origin": "form", "submit": True,
        "items": [item]})
    assert r.status_code == 201, r.text
    r = _post(maker, _csv(_row("2026-08-29"), "", _row("2026-09-15"), _row("2026-07-01")))
    assert r.status_code == 409 and r.json()["error"]["code"] == "duplicate_record"
    assert r.json()["error"]["params"]["rows"] == [
        {"row": 4, "field": "date", "code": "duplicate_record"}]
    assert _batches() == 1  # faqat oldingi form paketi


@pytest.mark.parametrize("member", [
    "xl/VbaProject.bin", "XL/vbaproject.BIN", "xl/externallinks/externalLink1.xml",
    "xl/activeX/activeX1.xml", "xl/embeddings/oleObject1.bin"])
def test_case_variant_dangerous_parts_rejected(member: str) -> None:
    with pytest.raises(ValidationFailed) as exc:
        imports.parse_upload(BALANCE, _patched(_valid(), {member: b"x"}), XLSX)
    assert exc.value.code == "unsupported_file"


def test_busy_gate_and_fail_closed(monkeypatch) -> None:  # type: ignore[no-untyped-def]
    import openpyxl

    from app.core.errors import AppError

    assert openpyxl.DEFUSEDXML is True
    assert imports._xlsx_gate.acquire(blocking=False)
    try:
        with pytest.raises(AppError) as exc:
            imports.parse_upload(BALANCE, _valid(), XLSX)
        assert exc.value.code == "import_busy" and exc.value.status_code == 429
    finally:
        imports._xlsx_gate.release()
    monkeypatch.setattr(openpyxl, "DEFUSEDXML", False)
    with pytest.raises(ValidationFailed) as exc2:
        imports.parse_upload(BALANCE, _valid(), XLSX)
    assert exc2.value.code == "unsupported_file"


def test_file_name_is_sanitised() -> None:
    raw = quote("C:\\fakepath\\a\u202e\u200bb\x85c\u2028.csv")
    assert imports.safe_name(raw) == "abc.csv"
    assert len(imports.safe_name("x" * 500) or "") == 200
    assert imports.safe_name("\u200b") is None


def test_import_route_declares_a_body_so_it_is_read_before_the_db_slot() -> None:
    from fastapi.routing import APIRoute

    app = create_app()
    routes = [r for r in app.routes if isinstance(r, APIRoute) and r.path.endswith("/import")]
    if routes:
        assert routes[0].body_field is not None
    else:  # lazy _IncludedRouter (fastapi 0.141): OpenAPI schema bor
        spec = app.openapi()["paths"]["/api/v1/data/entities/{entity_id}/import"]["post"]
        assert "requestBody" in spec


def test_legit_5000_row_xlsx_fits_the_caps(tmp_path) -> None:  # type: ignore[no-untyped-def]
    import os
    import subprocess
    import sys

    code = branch_code(1)
    rows = [[code, BANK, f"2026-{i % 12 + 1:02d}-{i % 28 + 1:02d}", f"{i}.25"]
            for i in range(5000)]
    path = tmp_path / "f.xlsx"
    path.write_bytes(_xlsx(rows))
    res = subprocess.run([sys.executable, "-c", PARSE_SCRIPT, str(path)], capture_output=True,
                         text=True, check=True, cwd="/app",
                         env={**os.environ, "PYTHONPATH": "/app"})
    out = json.loads(res.stdout.strip().splitlines()[-1])
    print("legit_5000", path.stat().st_size, out)  # noqa: T201
    assert out["out"] == "ok" and out["secs"] < 5 and out["mib"] < 150


# --- follow-ups: русские заголовки, один проход ошибок -----------------------------------------


def test_russian_headers_accepted_alone_and_mixed(client_for, make_user: MakeUser) -> None:  # type: ignore[no-untyped-def]
    maker = client_for(make_user("spec_kazna"))
    assert _post(maker, _csv(_row(), head=";".join(RU_HEAD))).status_code == 201
    mixed = "Филиал (код);bank_code;дата;BALANCE"  # регистр и смесь не важны
    assert _post(maker, _csv(_row("2026-08-29"), head=mixed)).status_code == 201
    dup = "Филиал (код);branch_code;date;balance"  # одно поле дважды -- ошибка шапки
    r = _post(maker, _csv(_row("2026-07-29"), head=dup))
    assert r.status_code == 422 and _codes(r) == [(1, "", "bad_header")]


def test_format_and_reference_errors_in_one_pass(client_for, make_user: MakeUser) -> None:  # type: ignore[no-untyped-def]
    maker = client_for(make_user("spec_kazna"))
    bad_ref = f"{branch_code(1)};BA-nonexistent;2026-08-29;1"
    bad_fmt = _row("2026-07-28", "abc")
    r = _post(maker, _csv(_row(), bad_fmt, bad_ref))
    assert r.status_code == 422 and r.json()["error"]["code"] == "import_invalid"
    assert _codes(r) == [(3, "balance", "invalid_value"), (4, "bank_code", "reference_not_found")]
    assert _batches() == 0
